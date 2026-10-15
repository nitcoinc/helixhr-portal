"""Plan 2026-10-04-001 U7: the monthly project timesheet (flagship) -- the
task x day grid, the detail, the pending-hours footnote, the exports.

Each test method uses its own past month: ERPNext's overlap check spans
every Timesheet of the shared fixture employees, and `IntegrationTestCase`
rolls back once per class, not per method.
"""

import io

import frappe
from frappe.tests import IntegrationTestCase

from helixhr import reports
from helixhr.api import run_report
from helixhr.tests.test_api_reports import _ensure_task, _make_timesheet
from helixhr.tests.utils import (
	ensure_baseline_company,
	ensure_test_company,
	make_test_delivery_manager,
	make_test_employee_and_manager,
	make_test_hr_manager_employee,
	make_test_project,
	set_report_access,
)

KEY = "project_timesheet"
HOLIDAY_LIST = "_Test U7 Holidays"


def _data(result):
	return [row for row in result["rows"] if row["_kind"] == "row"]


class TestProjectTimesheet(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.company = ensure_test_company()
		self.employee, _user, _, _ = make_test_employee_and_manager()
		self.hr_employee, self.hr_user = make_test_hr_manager_employee()
		self.dm_employee, self.dm_user = make_test_delivery_manager()
		set_report_access(KEY, dm_run=1, dm_export=1)
		self.project = make_test_project(self.company, "_Test U7 Flagship", members=[self.dm_user])
		self.other_project = make_test_project(self.company, "_Test U7 Not Member")
		self.task_a = _ensure_task(self.project, "_Test U7 Build")
		self.task_b = _ensure_task(self.project, "_Test U7 Review")

	def tearDown(self):
		frappe.set_user("Administrator")

	def _run(self, month, user=None, **filters):
		frappe.set_user(user or self.hr_user)
		try:
			return run_report(KEY, filters={"project": self.project, "month": month, **filters})
		finally:
			frappe.set_user("Administrator")

	def test_grid_cells_row_and_column_totals_match_the_detail(self):
		month = "2019-01"
		_make_timesheet(
			self.company,
			self.employee,
			"Emp",
			[
				{"date": "2019-01-07", "project": self.project, "task": self.task_a, "hours": 2},
				{"date": "2019-01-08", "project": self.project, "task": self.task_b, "hours": 1.5},
				{"date": "2019-01-09", "project": self.project, "task": self.task_a, "hours": 3},
			],
		)
		_make_timesheet(
			self.company,
			self.hr_employee,
			"HR",
			[
				{"date": "2019-01-07", "project": self.project, "task": self.task_a, "hours": 1},
				{"date": "2019-01-09", "project": self.project, "task": self.task_b, "hours": 0.25},
			],
		)
		result = self._run(month)
		grid = result["extra"]["grid"]
		self.assertEqual(len(grid["days"]), 31)
		by_task = {row["task"]: row for row in grid["rows"]}
		build, review = by_task["_Test U7 Build"], by_task["_Test U7 Review"]
		self.assertEqual(build["cells"][6], 3)  # 7th: 2 + 1
		self.assertEqual(build["cells"][8], 3)
		self.assertEqual(build["total"], 6)
		self.assertEqual(review["cells"][7], 1.5)
		self.assertEqual(review["total"], 1.75)
		self.assertEqual(grid["day_totals"][6], 3)
		self.assertEqual(grid["day_totals"][8], 3.25)
		self.assertIsNone(grid["day_totals"][0])
		self.assertEqual(grid["grand_total"], 7.75)
		self.assertEqual(result["totals"]["hours"], grid["grand_total"])
		# Detail grouped by date with daily subtotals.
		self.assertEqual(result["groups_applied"], ["date"])
		self.assertEqual(sum(1 for row in result["rows"] if row["_kind"] == "subtotal"), 3)
		self.assertEqual(result["extra"]["pending_hours"], 0)

	def test_pending_hours_are_excluded_and_counted_and_drafts_are_not(self):
		month = "2019-03"
		_make_timesheet(
			self.company, self.employee, "E", [{"date": "2019-03-04", "project": self.project, "hours": 2}]
		)
		_make_timesheet(
			self.company,
			self.hr_employee,
			"H",
			[{"date": "2019-03-05", "project": self.project, "hours": 4}],
			state="Pending Approval",
		)
		_make_timesheet(
			self.company,
			self.dm_employee,
			"D",
			[{"date": "2019-03-06", "project": self.project, "hours": 8}],
			state="Draft",
		)
		result = self._run(month)
		self.assertEqual(result["totals"]["hours"], 2)
		self.assertEqual(result["extra"]["grid"]["grand_total"], 2)
		self.assertEqual(result["extra"]["pending_hours"], 4)

	def test_a_row_without_a_task_is_grouped_under_no_task(self):
		_make_timesheet(
			self.company, self.employee, "E", [{"date": "2019-05-06", "project": self.project, "hours": 1}]
		)
		rows = self._run("2019-05")["extra"]["grid"]["rows"]
		self.assertEqual([row["task"] for row in rows], [reports.NO_TASK])

	def test_month_lengths_and_month_edges(self):
		for month, days in (("2019-02", 28), ("2019-04", 30), ("2019-07", 31)):
			self.assertEqual(len(self._run(month)["extra"]["grid"]["days"]), days, month)
		_make_timesheet(
			self.company,
			self.employee,
			"E",
			[
				{"date": "2019-06-30", "project": self.project, "hours": 2},
				{"date": "2019-07-01", "project": self.project, "hours": 5},
			],
		)
		june = self._run("2019-06")
		self.assertEqual(june["totals"]["hours"], 2)
		self.assertEqual(june["extra"]["grid"]["rows"][0]["cells"][29], 2)
		self.assertEqual(self._run("2019-07")["totals"]["hours"], 5)

	def test_weekends_and_company_holidays_are_marked(self):
		if not frappe.db.exists("Holiday List", HOLIDAY_LIST):
			frappe.get_doc(
				{
					"doctype": "Holiday List",
					"holiday_list_name": HOLIDAY_LIST,
					"from_date": "2019-01-01",
					"to_date": "2019-12-31",
					"holidays": [
						{"holiday_date": "2019-08-15", "description": "Founders' Day", "weekly_off": 0},
						{"holiday_date": "2019-08-17", "description": "Weekly Off", "weekly_off": 1},
					],
				}
			).insert(ignore_permissions=True)
		frappe.db.set_value("Company", self.company, "default_holiday_list", HOLIDAY_LIST)
		grid = self._run("2019-08")["extra"]["grid"]
		days = {d["day"]: d for d in grid["days"]}
		self.assertEqual(days[15]["holiday"], "Founders' Day")
		self.assertIsNone(days[17]["holiday"])  # weekly off: shaded, not "H"
		self.assertTrue(days[17]["weekend"])  # 2019-08-17 is a Saturday
		self.assertFalse(days[16]["weekend"])
		self.assertEqual(grid["holidays"], [{"date": "2019-08-15", "description": "Founders' Day"}])

	def test_billable_basis_sums_billing_hours(self):
		_make_timesheet(
			self.company,
			self.employee,
			"E",
			[
				{"date": "2019-09-02", "project": self.project, "hours": 3, "is_billable": 1},
				{"date": "2019-09-03", "project": self.project, "hours": 2, "is_billable": 0},
			],
		)
		all_hours = self._run("2019-09")["extra"]["grid"]
		billable = self._run("2019-09", basis="Billable hours")["extra"]["grid"]
		self.assertEqual(all_hours["grand_total"], 5)
		self.assertEqual(billable["grand_total"], 3)
		self.assertEqual(billable["field"], "billing_hours")

	def test_delivery_manager_member_project_only(self):
		_make_timesheet(
			self.company, self.employee, "E", [{"date": "2019-10-01", "project": self.project, "hours": 1}]
		)
		self.assertEqual(self._run("2019-10", user=self.dm_user)["totals"]["hours"], 1)
		frappe.set_user(self.dm_user)
		try:
			refused = run_report(KEY, filters={"project": self.other_project, "month": "2019-10"})
		finally:
			frappe.set_user("Administrator")
		self.assertEqual(refused["filters_removed"], ["project"])
		self.assertEqual(_data(refused), [])
		self.assertIsNone(refused["extra"])

	def test_another_companys_project_is_out_of_scope_for_hr(self):
		other = make_test_project(ensure_baseline_company(), "_Test U7 OtherCo")
		frappe.set_user(self.hr_user)
		try:
			result = run_report(KEY, filters={"project": other, "month": "2019-10"})
		finally:
			frappe.set_user("Administrator")
		self.assertEqual(result["filters_removed"], ["project"])

	def test_no_money_in_columns_grid_or_template_context(self):
		_make_timesheet(
			self.company,
			self.employee,
			"E",
			[{"date": "2019-11-04", "project": self.project, "hours": 1, "is_billable": 1}],
		)
		result = self._run("2019-11")
		money = ("rate", "amount", "cost")
		for column in result["columns"]:
			self.assertFalse(any(word in column["fieldname"] for word in money), column)
		self.assertFalse(any(word in key for key in result["extra"] for word in money))
		for row in result["rows"]:
			self.assertFalse(any(word in key for key in row for word in money), row)

	def test_a_script_note_is_escaped_in_the_pdf(self):
		sheet = _make_timesheet(
			self.company, self.employee, "E", [{"date": "2019-12-02", "project": self.project, "hours": 1}]
		)
		# Written raw: Frappe sanitises the field on save, but the template
		# must not rely on that.
		frappe.db.set_value(
			"Timesheet Detail", {"parent": sheet.name}, "description", "<script>alert(1)</script>"
		)
		html = self._pdf_html("2019-12")
		self.assertNotIn("<script>alert", html)
		self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", html)
		self.assertIn("Customer approval", html)
		self.assertIn("Prepared by", html)

	def test_empty_month_says_no_approved_hours_and_footnote_shows_pending(self):
		_make_timesheet(
			self.company,
			self.employee,
			"E",
			[{"date": "2018-01-02", "project": self.project, "hours": 6}],
			state="Pending HR",
		)
		result = self._run("2018-01")
		self.assertEqual(_data(result), [])
		self.assertEqual(result["extra"]["grid"]["rows"], [])
		html = self._pdf_html("2018-01")
		self.assertIn("No approved hours in this period.", html)
		self.assertIn("6 h awaiting approval are not included.", html)

	def _pdf_html(self, month):
		entry = reports.get_entry(KEY)
		scope = {"kind": "company", "company": self.company}
		result = reports.run(KEY, scope, {"project": self.project, "month": month})
		meta = reports.export_meta(entry, result, scope)
		columns = reports.visible_columns(result["columns"], [])
		return reports.render_pdf_html(
			entry,
			meta,
			columns,
			result["shaped"]["rows"],
			{"header": None, "footer": None, "logo": None, "company": self.company},
			result["extra"],
		)

	def test_excel_has_grid_and_detail_sheets_and_csv_is_detail_only(self):
		from openpyxl import load_workbook

		_make_timesheet(
			self.company,
			self.employee,
			"E",
			[
				{"date": "2018-02-05", "project": self.project, "task": self.task_a, "hours": 2},
				{"date": "2018-02-06", "project": self.project, "task": self.task_b, "hours": 1},
			],
		)
		entry = reports.get_entry(KEY)
		scope = {"kind": "company", "company": self.company}
		result = reports.run(KEY, scope, {"project": self.project, "month": "2018-02"})
		content, filename, _type = reports.build_export(entry, result, "xlsx", scope)
		wb = load_workbook(io.BytesIO(content))
		self.assertEqual(wb.sheetnames, ["Grid", "Detail"])
		grid_rows = [list(row) for row in wb["Grid"].iter_rows(values_only=True)]
		header = next(i for i, row in enumerate(grid_rows) if row and row[0] == "Task")
		self.assertEqual(len(grid_rows[header]), 1 + 28 + 1)
		self.assertEqual(grid_rows[-1][0], "Total")
		self.assertEqual(grid_rows[-1][-1], 3)
		self.assertIsInstance(grid_rows[header + 1][5], float | int)  # Feb 5th cell is numeric
		self.assertTrue(filename.endswith(".xlsx"))

		csv, _name, _type = reports.build_export(entry, result, "csv", scope)
		lines = csv[3:].decode().strip().splitlines()
		self.assertEqual(len(lines), 1 + 2)  # header + two detail rows, no subtotals/total
		self.assertNotIn("Total", csv.decode())


# --- Plan 2026-10-15-001 U2: combined and logged-by project reports ----------


class TestCrossCompanyProjectReports(IntegrationTestCase):
	"""The Nitco case: an Inc project, 5h logged by Inc staff and 3h by
	Outsourcing staff (`ensure_cross_company_project_fixture`)."""

	OUT_HR = "xco-out-hr@helixhr.test"

	def setUp(self):
		from helixhr.tests.utils import (
			GRANT_HOLDER_USER,
			_make_role_user,
			ensure_cross_company_project_fixture,
			make_test_company_grant,
			make_test_user,
		)

		frappe.set_user("Administrator")
		self.fx = ensure_cross_company_project_fixture()
		_hr_employee, self.inc_hr = make_test_hr_manager_employee()
		_make_role_user(self.OUT_HR, "HR Manager", self.fx.other, create_user_permission=0)
		self.holder = GRANT_HOLDER_USER
		make_test_user(self.holder, self.fx.company)
		make_test_company_grant(self.holder, [self.fx.company, self.fx.other], allow_export=1)

	def tearDown(self):
		frappe.set_user("Administrator")

	def _as(self, user, fn, *args, **kwargs):
		frappe.set_user(user)
		try:
			return fn(*args, **kwargs)
		finally:
			frappe.set_user("Administrator")

	def _run(self, user, key=KEY, **filters):
		filters = {"project": self.fx.project, "month": "2016-04", **filters} if key == KEY else filters
		return self._as(user, run_report, key, filters=filters)

	def _options(self, user, filter_name, context=None, key="hours_by_project"):
		from helixhr.api import search_report_options

		return {
			row["value"]
			for row in self._as(user, search_report_options, key, filter_name, context=context or {})
		}

	def test_combined_run_totals_both_companies_with_a_company_column(self):
		result = self._run(self.holder)
		self.assertEqual(result["filters"]["company"], reports.SENTINEL_ALL)
		self.assertIn("company", [column["fieldname"] for column in result["columns"]])
		hours = {}
		for row in _data(result):
			hours[row["company"]] = hours.get(row["company"], 0) + row["hours"]
		self.assertEqual(hours, {self.fx.company: 5, self.fx.other: 3})
		self.assertEqual(result["company_counts"], {self.fx.company: 1, self.fx.other: 1})
		self.assertEqual(result["extra"]["grid"]["grand_total"], 8)

	def test_inc_hr_without_a_grant_sees_inc_hours_only(self):
		result = self._run(self.inc_hr)
		self.assertEqual(sum(row["hours"] for row in _data(result)), 5)
		self.assertNotIn("company", [column["fieldname"] for column in result["columns"]])
		self.assertIsNone(result["company_counts"])

	def test_outsourcing_hr_picks_the_inc_project_and_sees_outsourcing_hours_and_holidays(self):
		"""Covers AE3."""
		from unittest.mock import patch

		with patch.object(reports, "_company_holidays", return_value={}) as holidays:
			result = self._run(self.OUT_HR)
		self.assertEqual(result["filters_removed"], [])
		self.assertEqual(sum(row["hours"] for row in _data(result)), 3)
		self.assertEqual(holidays.call_args.args[0], self.fx.other)

	def test_combined_run_labels_project_company_holidays(self):
		from unittest.mock import patch

		day = frappe.utils.getdate("2016-04-14")
		with patch.object(reports, "_company_holidays", return_value={day: "Fixture day"}) as holidays:
			result = self._run(self.holder)
		self.assertEqual(holidays.call_args.args[0], self.fx.company)
		self.assertEqual(
			result["extra"]["grid"]["holidays"], [{"date": "2016-04-14", "description": f"Fixture day ({self.fx.company})"}]
		)

	def test_combined_hours_by_project_groups_both_employees(self):
		result = self._as(
			self.holder,
			run_report,
			"hours_by_project",
			filters={"project": self.fx.project, "from_date": "2016-04-01", "to_date": "2016-04-30"},
			group_by=["project"],
		)
		rows = _data(result)
		self.assertEqual({row["employee"] for row in rows}, {self.fx.inc, self.fx.out})
		self.assertEqual({row["company"] for row in rows}, {self.fx.company, self.fx.other})

	def test_outsourcing_pickers_list_logged_projects_tasks_and_employees_only(self):
		projects = self._options(self.OUT_HR, "project", key=KEY)
		self.assertIn(self.fx.project, projects)
		self.assertNotIn(self.fx.idle, projects)
		tasks = self._options(self.OUT_HR, "task", {"project": self.fx.project})
		self.assertEqual(tasks, {self.fx.out_task})
		# Inc HR owns the project: every task on it.
		self.assertEqual(
			self._options(self.inc_hr, "task", {"project": self.fx.project}) & {self.fx.inc_task, self.fx.out_task},
			{self.fx.inc_task, self.fx.out_task},
		)
		employees = self._as(
			self.OUT_HR,
			reports.search_options,
			reports.get_entry("hours_by_project"),
			"employee",
			{"kind": "company", "company": self.fx.other},
			context={"project": self.fx.project},
		)
		self.assertEqual({row["value"] for row in employees}, {self.fx.out})

	def test_out_of_reach_project_and_task_are_removed(self):
		result = self._run(self.OUT_HR, project=self.fx.idle)
		self.assertEqual(result["filters_removed"], ["project"])
		removed = self._as(
			self.OUT_HR,
			run_report,
			"hours_by_project",
			filters={"project": self.fx.project, "task": self.fx.inc_task},
		)
		self.assertEqual(removed["filters_removed"], ["task"])

	def test_pending_hours_obey_the_same_scope(self):
		_make_timesheet(
			self.fx.company,
			self.fx.inc,
			"XCO Inc",
			[{"date": "2016-04-20", "project": self.fx.project, "task": self.fx.inc_task, "hours": 2}],
			state="Pending Approval",
		)
		self.assertEqual(self._run(self.OUT_HR)["extra"]["pending_hours"], 0)
		self.assertEqual(self._run(self.inc_hr)["extra"]["pending_hours"], 2)

	def test_a_third_company_never_appears_in_a_combined_run(self):
		from helixhr.tests.utils import COMPANY_ACCESS_THIRD_COMPANY, create_test_company, make_test_user

		third = create_test_company(COMPANY_ACCESS_THIRD_COMPANY, "TCAC")
		employee = make_test_user("xco-third@helixhr.test", third)
		_make_timesheet(
			third,
			employee,
			"XCO Third",
			[{"date": "2016-04-21", "project": self.fx.project, "task": self.fx.inc_task, "hours": 4}],
		)
		result = self._run(self.holder)
		self.assertEqual(sum(row["hours"] for row in _data(result)), 8)
		self.assertNotIn(third, {row["company"] for row in _data(result)})

	def test_projects_page_is_not_widened(self):
		from helixhr.api import search_projects

		names = {row["name"] for row in self._as(self.OUT_HR, search_projects)["projects"]}
		self.assertNotIn(self.fx.project, names)

	def test_an_unknown_scope_kind_raises(self):
		with self.assertRaises(frappe.ValidationError):
			reports._timesheet_scope_conditions({"kind": "companies", "companies": [self.fx.company]}, {})
