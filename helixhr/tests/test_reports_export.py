"""Plan 2026-10-04-001 U5: inline CSV / Excel / PDF export, the audit row,
the one-time download token and the export log."""

import datetime
import io
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase

from helixhr import reports
from helixhr.api import (
	download_export,
	download_report_export,
	get_export_log,
	list_my_exports,
	request_export,
)
from helixhr.helixhr.doctype.helixhr_report_export.helixhr_report_export import HelixHRReportExport
from helixhr.tests.utils import (
	ensure_test_company,
	make_test_hr_manager_employee,
	make_test_hr_user,
	make_test_report_manager,
	make_test_user,
	set_report_access,
)

KEY = "employee_directory"
HOURS_COLUMNS = [
	{"fieldname": "date", "label": "Date", "fieldtype": "Date"},
	{"fieldname": "employee", "label": "Employee", "fieldtype": "Link"},
	{"fieldname": "note", "label": "Note", "fieldtype": "Data"},
	{"fieldname": "hours", "label": "Hours", "fieldtype": "Float"},
]
HOURS_ROWS = [
	{"date": datetime.date(2026, 9, 1), "employee": "E1", "note": '=HYPERLINK("x")', "hours": 1.5},
	{"date": datetime.date(2026, 9, 2), "employee": "E1", "note": "fine", "hours": -0.25},
	{"date": "2026-09-03", "employee": "E2", "note": "", "hours": 2},
]
# 1x1 transparent PNG.
PNG = bytes.fromhex(
	"89504e470d0a1a0a0000000d4948445200000001000000010806000000"
	"1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
)


def _as(user, fn, *args, **kwargs):
	frappe.set_user(user)
	try:
		return fn(*args, **kwargs)
	finally:
		frappe.set_user("Administrator")


class TestExportFormats(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.entry = reports.get_entry(KEY)
		self.shaped = reports.shape(HOURS_COLUMNS, HOURS_ROWS, ["employee"])

	def test_csv_has_bom_escapes_formulas_and_keeps_numbers_numeric(self):
		content = reports.to_csv(HOURS_COLUMNS, self.shaped["rows"])
		self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
		text = content[3:].decode()
		self.assertIn("'=HYPERLINK", text)
		self.assertNotIn(",=HYPERLINK", text)
		self.assertIn(",-0.25\r\n", text)
		self.assertIn(",1.5\r\n", text)
		self.assertIn("Subtotal: E1", text)
		self.assertTrue(text.rstrip().endswith("Total,,,3.25"))

	def test_xlsx_totals_match_the_shaper_and_dates_are_dates(self):
		from openpyxl import load_workbook

		meta = {
			"title": "Hours",
			"company": "Co",
			"period": "",
			"filters": [],
			"group_by": ["Employee"],
			"generated_by": "Administrator",
			"generated_at": "2026-10-04 10:00 UTC",
		}
		result = {"shaped": self.shaped, "columns": HOURS_COLUMNS}
		ws = load_workbook(io.BytesIO(reports.to_xlsx(self.entry, meta, HOURS_COLUMNS, result))).active
		rows = [list(row) for row in ws.iter_rows(values_only=True)]
		header = rows.index(["Date", "Employee", "Note", "Hours"])
		body = rows[header + 1 :]
		self.assertEqual(body[-1][0], "Total")
		self.assertEqual(body[-1][3], self.shaped["totals"]["hours"])
		self.assertIsInstance(body[0][0], datetime.datetime)
		self.assertIsInstance(body[2][0], str)  # the subtotal label
		self.assertTrue(ws.cell(row=header + 2 + 2, column=1).font.b)  # subtotal bold
		self.assertEqual(body[0][2], '=HYPERLINK("x")')  # literal text, not a formula
		self.assertEqual(ws.cell(row=header + 2, column=3).data_type, "s")

	def test_filename_is_ascii_with_key_period_and_stamp(self):
		name = reports.export_filename(
			self.entry,
			{"from_date": "2026-09-01", "to_date": "2026-09-30"},
			"xlsx",
			datetime.datetime(2026, 10, 4, 9, 5),
		)
		self.assertEqual(name, "helixhr_employee_directory_2026-09-01_2026-09-30_20261004T0905.xlsx")


class TestPdfLetterHead(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.company = ensure_test_company()
		self.entry = reports.get_entry(KEY)
		self.meta = {
			"title": "Employee directory",
			"company": self.company,
			"period": "",
			"filters": [("Status", "Active")],
			"group_by": [],
			"generated_by": "Administrator",
			"generated_at": "2026-10-04 10:00 UTC",
		}
		self.columns = [{"fieldname": "note", "label": "Note", "fieldtype": "Data"}]
		self.saved = frappe.db.get_value(
			"Company", self.company, ["default_letter_head", "company_logo"], as_dict=True
		)
		frappe.db.set_value("Company", self.company, {"default_letter_head": None, "company_logo": None})
		self.default_heads = frappe.get_all("Letter Head", filters={"is_default": 1}, pluck="name")
		for name in self.default_heads:
			frappe.db.set_value("Letter Head", name, "is_default", 0)

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.db.set_value("Company", self.company, self.saved)
		for name in self.default_heads:
			frappe.db.set_value("Letter Head", name, "is_default", 1)

	def _private_png(self, name):
		existing = frappe.db.get_value("File", {"file_name": name, "is_private": 1}, "file_url")
		if existing:
			return existing
		return (
			frappe.get_doc({"doctype": "File", "file_name": name, "is_private": 1, "content": PNG})
			.insert(ignore_permissions=True)
			.file_url
		)

	def _letter_head(self, content):
		name = "HelixHR Export Test Head"
		if frappe.db.exists("Letter Head", name):
			frappe.db.set_value("Letter Head", name, {"source": "HTML", "content": content, "disabled": 0})
		else:
			frappe.get_doc(
				{"doctype": "Letter Head", "letter_head_name": name, "source": "HTML", "content": content}
			).insert(ignore_permissions=True)
		frappe.db.set_value("Company", self.company, "default_letter_head", name)

	def _html(self, rows=()):
		return reports.render_pdf_html(self.entry, self.meta, self.columns, [*rows, {"_kind": "total"}])

	def test_company_letter_head_is_static_html_with_images_inlined(self):
		url = self._private_png("helixhr-export-lh.png")
		self._letter_head(
			f'<div class="lh">ACME Letter<img src="{url}"><script>x()</script></div>'
		)
		html = self._html()
		self.assertIn('id="header-html"', html)
		self.assertIn("ACME Letter", html)
		self.assertIn("data:image/png;base64,", html)
		self.assertNotIn(url, html)
		self.assertNotIn("<script>x()", html)

	def test_jinja_in_a_letter_head_falls_back_instead_of_printing_raw(self):
		self._letter_head("<div>{% if doc %}ACME{% endif %} {{ frappe.session.user }}</div>")
		frappe.db.set_value("Letter Head", "HelixHR Export Test Head", "footer", "<p>{{ company }}</p>")
		try:
			html = self._html()
		finally:
			frappe.db.set_value("Letter Head", "HelixHR Export Test Head", "footer", None)
		self.assertNotIn('id="header-html"', html)
		self.assertNotIn("{%", html)
		self.assertNotIn("{{", html)
		self.assertNotIn('class="lh-footer"', html)
		self.assertIn(self.company, html)

		lh = reports.resolve_letter_head(self.company)
		self.assertEqual((lh["header"], lh["footer"]), (None, None))

	def test_without_a_letter_head_logo_and_name_then_name_alone(self):
		html = self._html()
		self.assertNotIn('id="header-html"', html)
		self.assertIn(self.company, html)
		self.assertNotIn("<img", html)

		frappe.db.set_value(
			"Company", self.company, "company_logo", self._private_png("helixhr-export-logo.png")
		)
		self.assertIn('src="data:image/png;base64,', self._html())

	def test_cells_are_escaped(self):
		html = self._html([{"_kind": "row", "note": "<script>alert(1)</script>"}])
		self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", html)

	def test_pdf_bytes_and_page_numbering(self):
		options = reports.pdf_options(self.entry, self.meta, self.columns, False)
		self.assertIn("[page]", options["footer-right"])
		self.assertIn("[topage]", options["footer-right"])
		self.assertIn("Administrator", options["footer-left"])
		rows = [{"_kind": "row", "note": f"line {i}"} for i in range(120)]
		pdf = reports.to_pdf(self.entry, self.meta, self.columns, [*rows, {"_kind": "total"}])
		self.assertTrue(pdf.startswith(b"%PDF"))

	def test_wide_reports_go_landscape(self):
		wide = [{"fieldname": f"c{i}", "label": str(i)} for i in range(7)]
		self.assertEqual(reports.pdf_options(self.entry, self.meta, wide, False)["orientation"], "Landscape")
		self.assertEqual(
			reports.pdf_options(self.entry, self.meta, self.columns, False)["orientation"], "Portrait"
		)

	def test_a_missing_generator_is_one_plain_sentence(self):
		with patch("frappe.utils.pdf.get_pdf", side_effect=OSError("wkhtmltopdf: not found")):
			with patch("frappe.log_error") as log:
				with self.assertRaises(frappe.ValidationError) as caught:
					reports.to_pdf(self.entry, self.meta, self.columns, [{"_kind": "total"}])
		self.assertIn("PDF export isn't available", str(caught.exception))
		log.assert_called_once()


class TestRequestExport(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		self.company = ensure_test_company()
		make_test_user("export-colleague@helixhr.test", self.company)
		_, self.hr_user = make_test_hr_user()
		if not getattr(frappe.local, "response_headers", None):
			from werkzeug.datastructures import Headers

			frappe.local.response_headers = Headers()

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.local.response = frappe._dict({"docs": []})

	def test_run_without_export_is_refused(self):
		set_report_access(KEY, hr_user_run=1)
		with self.assertRaises(frappe.PermissionError):
			_as(self.hr_user, request_export, KEY, "csv")

	def test_granted_export_logs_a_row_and_downloads_once_for_its_owner(self):
		set_report_access(KEY, hr_user_run=1, hr_user_export=1)
		out = _as(self.hr_user, request_export, KEY, "csv", filters={"status": "Active"}, hidden=["gender"])
		self.assertRegex(out["filename"], r"^helixhr_employee_directory_\d{8}T\d{4}\.csv$")

		log = frappe.get_doc("HelixHR Report Export", out["export"])
		self.assertEqual(
			(log.owner, log.report_key, log.format, log.mode, log.company, log.row_count),
			(self.hr_user, KEY, "csv", "Inline", self.company, out["row_count"]),
		)
		self.assertEqual(frappe.parse_json(log.filters), {"status": "Active"})

		# Another user cannot redeem it, and that attempt does not burn it.
		with self.assertRaises(frappe.PermissionError):
			_as("Administrator", download_export, out["token"])

		_as(self.hr_user, download_export, out["token"])
		self.assertEqual(frappe.local.response.type, "download")
		self.assertTrue(frappe.local.response.filecontent.startswith(b"\xef\xbb\xbf"))
		self.assertNotIn(b"Gender", frappe.local.response.filecontent)
		self.assertEqual(frappe.local.response_headers["Cache-Control"], "no-store")
		self.assertTrue(
			frappe.local.response_headers["Content-Disposition"].startswith("attachment; filename*=")
		)

		with self.assertRaises(frappe.PermissionError):
			_as(self.hr_user, download_export, out["token"])

	def test_xlsx_and_pdf_render_through_the_endpoint(self):
		set_report_access(KEY, hr_user_run=1, hr_user_export=1)
		for fmt in ("xlsx", "pdf"):
			out = _as(self.hr_user, request_export, KEY, fmt)
			self.assertTrue(out["filename"].endswith(f".{fmt}"))

	def test_unknown_or_bad_format_is_refused(self):
		with self.assertRaises(frappe.PermissionError) as unknown:
			_as(self.hr_user, request_export, "not_a_report", "csv")
		with self.assertRaises(frappe.PermissionError) as deny_listed:
			_as(self.hr_user, request_export, "Timesheet Billing Summary", "csv")
		self.assertEqual(str(unknown.exception), str(deny_listed.exception))
		with self.assertRaises(frappe.ValidationError):
			_as("Administrator", request_export, KEY, "docx")

	def test_above_both_caps_says_narrow_the_filters(self):
		with (
			patch.dict(reports.INLINE_EXPORT_CAP, {"csv": 0}),
			patch.dict(reports.BACKGROUND_EXPORT_CAP, {"csv": 0}),
		):
			with self.assertRaises(frappe.ValidationError) as caught:
				_as("Administrator", request_export, KEY, "csv")
		self.assertIn("Narrow the filters", str(caught.exception))

	def test_export_log_is_hr_manager_only(self):
		_, report_manager = make_test_report_manager()
		set_report_access(KEY, hr_user_run=1, hr_user_export=1)
		out = _as(self.hr_user, request_export, KEY, "csv")
		for user in (self.hr_user, report_manager):
			with self.assertRaises(frappe.PermissionError):
				_as(user, get_export_log)

		_, hr_manager = make_test_hr_manager_employee()
		page = _as(hr_manager, get_export_log, page_length=200)
		row = next(row for row in page["rows"] if row.name == out["export"])
		self.assertEqual(row.owner, self.hr_user)
		self.assertNotIn("file", row)


class TestBackgroundExport(IntegrationTestCase):
	"""U13. The job commits, so every row, File and Notification Log it
	makes is named and deleted explicitly (runbook isolation notes)."""

	def setUp(self):
		frappe.set_user("Administrator")
		self.company = ensure_test_company()
		_, self.hr_user = make_test_hr_user()
		set_report_access(KEY, hr_user_run=1, hr_user_export=1)
		self._clear()
		self.addCleanup(self._clear)
		if not getattr(frappe.local, "response_headers", None):
			from werkzeug.datastructures import Headers

			frappe.local.response_headers = Headers()
		caps = patch.dict(reports.INLINE_EXPORT_CAP, {"csv": 0, "xlsx": 0, "pdf": 0})
		caps.start()
		self.addCleanup(caps.stop)

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.local.response = frappe._dict({"docs": []})

	def _clear(self):
		frappe.set_user("Administrator")
		names = frappe.get_all(
			"HelixHR Report Export", filters={"owner": self.hr_user, "mode": "Background"}, pluck="name"
		)
		for name in names:
			for file in frappe.get_all(
				"File",
				filters={"attached_to_doctype": "HelixHR Report Export", "attached_to_name": name},
				pluck="name",
			):
				frappe.delete_doc("File", file, ignore_permissions=True, force=True)
			frappe.db.delete("Notification Log", {"document_type": "HelixHR Report Export", "document_name": name})
			frappe.db.delete("HelixHR Report Export", {"name": name})
		frappe.db.commit()  # nosemgrep -- the job under test commits too

	def _request(self, fmt="csv"):
		with patch("frappe.enqueue") as enqueue:
			out = _as(self.hr_user, request_export, KEY, fmt, filters={"status": "Active"})
		return out, enqueue

	def _run_job(self, enqueue):
		kwargs = dict(enqueue.call_args.kwargs)
		for key in ("queue", "job_id", "deduplicate", "enqueue_after_commit"):
			kwargs.pop(key)
		reports.run_background_export(**kwargs)

	def test_queues_then_job_attaches_a_file_and_notifies(self):
		out, enqueue = self._request()
		self.assertEqual(out["status"], "Queued")
		self.assertEqual(enqueue.call_args.args, ("helixhr.reports.run_background_export",))
		self.assertTrue(enqueue.call_args.kwargs["enqueue_after_commit"])
		self.assertTrue(enqueue.call_args.kwargs["job_id"].startswith(f"report-export:{self.hr_user}:"))
		row = frappe.get_doc("HelixHR Report Export", out["export"])
		self.assertEqual((row.mode, row.status, row.owner), ("Background", "Queued", self.hr_user))

		self._run_job(enqueue)
		row.reload()
		self.assertEqual(row.status, "Ready")
		self.assertTrue(row.file.startswith("/private/files/"))
		self.assertTrue(
			frappe.db.exists(
				"Notification Log",
				{"for_user": self.hr_user, "document_type": "HelixHR Report Export", "document_name": row.name},
			)
		)
		listed = _as(self.hr_user, list_my_exports)
		self.assertEqual(next(r for r in listed if r.name == row.name).status, "Ready")

		_as(self.hr_user, download_report_export, row.name)
		self.assertTrue(frappe.local.response.filecontent.startswith(b"\xef\xbb\xbf"))
		self.assertEqual(frappe.local.response_headers["Cache-Control"], "no-store")

		# Another HR Manager: refused by the endpoint and by the File chain.
		_, hr_manager = make_test_hr_manager_employee()
		with self.assertRaises(frappe.PermissionError):
			_as(hr_manager, download_report_export, row.name)
		file = frappe.get_doc("File", {"file_url": row.file, "attached_to_name": row.name})
		self.assertFalse(frappe.has_permission("File", "read", file, user=hr_manager))
		self.assertTrue(frappe.has_permission("File", "read", file, user=self.hr_user))
		self.assertNotIn(row.name, [r.name for r in _as(hr_manager, list_my_exports)])

	def test_revoked_access_fails_with_no_file(self):
		out, enqueue = self._request()
		set_report_access(KEY, hr_user_run=1, hr_user_export=0)
		self._run_job(enqueue)
		row = frappe.get_doc("HelixHR Report Export", out["export"])
		self.assertEqual(row.status, "Failed")
		self.assertFalse(row.file)
		self.assertFalse(frappe.db.exists("File", {"attached_to_name": row.name}))
		self.assertIn("no longer have access", row.error)

	def test_identical_request_dedups_and_two_active_is_the_cap(self):
		first, _ = self._request("csv")
		again, enqueue = self._request("csv")
		self.assertEqual(again["export"], first["export"])
		enqueue.assert_not_called()
		self._request("xlsx")
		with self.assertRaises(frappe.ValidationError):
			self._request("pdf")

	def test_above_the_background_cap_is_refused(self):
		with patch.dict(reports.BACKGROUND_EXPORT_CAP, {"csv": 0}):
			with self.assertRaises(frappe.ValidationError) as caught:
				self._request()
		self.assertIn("Narrow the filters", str(caught.exception))


class TestExportRetention(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")

	def _row(self, days_old, with_file=False):
		doc = frappe.get_doc(
			{"doctype": "HelixHR Report Export", "report_key": KEY, "format": "csv", "row_count": 1}
		).insert(ignore_permissions=True)
		if with_file:
			file = frappe.get_doc(
				{
					"doctype": "File",
					"file_name": f"{doc.name}.csv",
					"is_private": 1,
					"content": b"a,b",
					"attached_to_doctype": "HelixHR Report Export",
					"attached_to_name": doc.name,
				}
			).insert(ignore_permissions=True)
			frappe.db.set_value("HelixHR Report Export", doc.name, "file", file.file_url)
		old = frappe.utils.add_days(frappe.utils.now_datetime(), -days_old)
		frappe.db.set_value("HelixHR Report Export", doc.name, "creation", old, update_modified=False)
		return doc.name

	def test_files_expire_after_seven_days_and_rows_after_the_retention(self):
		fresh_file = self._row(6, with_file=True)
		stale_file = self._row(8, with_file=True)
		ancient = self._row(400)
		HelixHRReportExport.clear_old_logs(365)

		self.assertTrue(frappe.db.get_value("HelixHR Report Export", fresh_file, "file"))
		self.assertEqual(
			frappe.db.get_value("HelixHR Report Export", stale_file, ["file", "status"]), (None, "Expired")
		)
		self.assertFalse(frappe.db.exists("File", {"attached_to_name": stale_file}))
		self.assertFalse(frappe.db.exists("HelixHR Report Export", ancient))
		for name in (fresh_file, stale_file):
			frappe.delete_doc("HelixHR Report Export", name, ignore_permissions=True, force=True)
